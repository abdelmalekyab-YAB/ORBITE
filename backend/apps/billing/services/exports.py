"""
Exports (SPEC §8) : dossier comptable (Excel + ZIP de justificatifs) et rapport PDF direction.

Dépendances : openpyxl, reportlab.
"""

from __future__ import annotations

import contextlib
import io
import re
import unicodedata
import zipfile
from decimal import Decimal

from .billing import BillingEngine, month_range

MONTHS_FR = [
    "janvier",
    "février",
    "mars",
    "avril",
    "mai",
    "juin",
    "juillet",
    "août",
    "septembre",
    "octobre",
    "novembre",
    "décembre",
]


def month_label(m: str, short=False) -> str:
    y, mo = m[:4], int(m[5:7])
    name = MONTHS_FR[mo - 1]
    return f"{name[:4]}. {y[2:]}" if short else f"{name} {y}"


def eur(v, dec=2) -> str:
    v = Decimal(v or 0)
    s = f"{v:,.{dec}f}".replace(",", " ").replace(".", ",")
    return s + " €"


def safe(t: str) -> str:
    t = unicodedata.normalize("NFD", str(t or "")).encode("ascii", "ignore").decode()
    t = re.sub(r"[^\w\- ]+", " ", t).strip()
    return re.sub(r"\s+", "_", t)[:60] or "document"


# ---------------------------------------------------------------------------
# Données d'un export comptable
# ---------------------------------------------------------------------------
def accounting_data(e: BillingEngine, start: str, end: str, client_id: str | None, with_expenses: bool) -> dict:
    from apps.billing import models as M

    if end < start:
        start, end = end, start
    months = month_range(start, end)
    projs = [p for p in e.s.projects if not client_id or p.client_id == client_id]
    sales = []
    for m in months:
        for p in projs:
            c = e.client(p.client_id)
            if p.is_fixed and not p.is_proposal and p.active_in(m) and e.fixed_billed(p, m) and e.fixed_value(p, m):
                o = e.occurrence(p, m)
                doc = M.FixedOccurrence.objects.filter(project_id=p.id, month=m).first()
                sales.append(
                    {
                        "month": m,
                        "client": c.name if c else "—",
                        "project": p.name,
                        "label": "Montant fixe mensuel",
                        "amount": e.fixed_value(p, m),
                        "pay": e._pay_state(o.payment_status, o.due_date) if o else "wait",
                        "due": o.due_date if o else None,
                        "paid_on": o.paid_on if o else None,
                        "file": doc.document if (doc and doc.document) else None,
                        "file_name": doc.document_name if doc else "",
                    }
                )
            for i in e.invoices_of(p):
                if i.service_month == m and e.is_billed(i):
                    inv = M.Invoice.objects.filter(pk=i.id).first()
                    sales.append(
                        {
                            "month": m,
                            "client": c.name if c else "—",
                            "project": p.name,
                            "label": i.label or "Facture",
                            "amount": Decimal(i.amount),
                            "pay": e._pay_state(i.payment_status, i.due_date),
                            "due": i.due_date,
                            "paid_on": i.paid_on,
                            "file": inv.document if (inv and inv.document) else None,
                            "file_name": i.document_name,
                        }
                    )
    exps = []
    if with_expenses:
        pids = {p.id for p in projs}
        for x in M.Expense.objects.select_related("project").all():
            m = x.date.strftime("%Y-%m")
            if m in months and (not client_id or (x.project_id and str(x.project_id) in pids)):
                exps.append(x)
    client_name = e.client(client_id).name if client_id and e.client(client_id) else ""
    return {"start": start, "end": end, "months": months, "sales": sales, "expenses": exps, "client_name": client_name}


def build_workbook(d: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Synthèse"
    title = (
        f"DIGITALIA · Rapport {month_label(d['start'])}"
        + (f" – {month_label(d['end'])}" if d["start"] != d["end"] else "")
        + (f" · {d['client_name']}" if d["client_name"] else "")
    )
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    ws.append(["Client"] + [month_label(m, True) for m in d["months"]] + ["Total"])
    clients = sorted({s["client"] for s in d["sales"]})
    for c in clients:
        vals = [
            sum((s["amount"] for s in d["sales"] if s["client"] == c and s["month"] == m), Decimal(0))
            for m in d["months"]
        ]
        ws.append([c] + [float(v) for v in vals] + [float(sum(vals))])
    tot = [sum((s["amount"] for s in d["sales"] if s["month"] == m), Decimal(0)) for m in d["months"]]
    ws.append(["Total ventes"] + [float(v) for v in tot] + [float(sum(tot))])
    if d["expenses"]:
        dep = [
            sum((x.amount_ht for x in d["expenses"] if x.date.strftime("%Y-%m") == m), Decimal(0)) for m in d["months"]
        ]
        ws.append(["Dépenses HT"] + [float(v) for v in dep] + [float(sum(dep))])
        ws.append(["Résultat"] + [float(a - b) for a, b in zip(tot, dep, strict=False)] + [float(sum(tot) - sum(dep))])
    ws.column_dimensions["A"].width = 28
    ws2 = wb.create_sheet("Ventes")
    ws2.append(["Mois", "Client", "Projet", "Libellé", "Montant HT", "Paiement", "Échéance", "Payé le", "Justificatif"])
    PAY = {"paid": "Payé", "late": "En retard", "due": "En attente", "wait": "En attente"}
    for s in d["sales"]:
        ws2.append(
            [
                s["month"],
                s["client"],
                s["project"],
                s["label"],
                float(s["amount"]),
                PAY[s["pay"]],
                s["due"].isoformat() if s["due"] else "",
                s["paid_on"].isoformat() if s["paid_on"] else "",
                s["file_name"] or "MANQUANT",
            ]
        )
    for col, w in zip("ABCDEFGHI", [9, 18, 32, 48, 12, 12, 11, 11, 34], strict=False):
        ws2.column_dimensions[col].width = w
    if d["expenses"]:
        ws3 = wb.create_sheet("Dépenses")
        ws3.append(["Date", "Fournisseur", "Catégorie", "Description", "Projet", "HT", "TVA", "TTC", "Justificatif"])
        for x in d["expenses"]:
            ws3.append(
                [
                    x.date.isoformat(),
                    x.supplier,
                    x.category,
                    x.label,
                    x.project.name if x.project else "Frais généraux",
                    float(x.amount_ht),
                    float(x.vat),
                    float(x.amount_ht + x.vat),
                    x.receipt_name or "MANQUANT",
                ]
            )
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def accounting_package(e, start, end, client_id, with_expenses, fmt="zip"):
    d = accounting_data(e, start or e.now, end or e.now, client_id or None, with_expenses)
    base = (
        f"DIGITALIA_compta_{d['start']}"
        + (f"_{d['end']}" if d["start"] != d["end"] else "")
        + (f"_{safe(d['client_name'])}" if d["client_name"] else "")
    )
    xlsx = build_workbook(d)
    if fmt == "xlsx":
        return xlsx, base + ".xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    buf = io.BytesIO()
    used, missing = set(), []

    def uniq(path):
        p, i = path, 2
        while p in used:
            stem, ext = path.rsplit(".", 1)
            p = f"{stem}_{i}.{ext}"
            i += 1
        used.add(p)
        return p

    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr(base + ".xlsx", xlsx)
        for s in d["sales"]:
            if s["file"]:
                ext = (s["file_name"].rsplit(".", 1)[-1] if "." in s["file_name"] else "pdf").lower()
                with s["file"].open("rb") as fh:
                    z.writestr(
                        uniq(
                            f"Factures/{safe(s['client'])}/{s['month']}_{safe(s['project'])}_{safe(s['label'])}.{ext}"
                        ),
                        fh.read(),
                    )
            else:
                missing.append(
                    f"- Vente {s['month']} · {s['client']} · {s['project']} · {s['label']} · {eur(s['amount'])}"
                )
        for x in d["expenses"]:
            if x.receipt:
                ext = (x.receipt_name.rsplit(".", 1)[-1] if "." in x.receipt_name else "pdf").lower()
                with x.receipt.open("rb") as fh:
                    z.writestr(
                        uniq(
                            f"Depenses/{x.date:%Y-%m}/{x.date}_{safe(x.supplier)}_{safe(x.label or x.category)}.{ext}"
                        ),
                        fh.read(),
                    )
            else:
                missing.append(f"- Dépense {x.date} · {x.supplier} · {eur(x.amount_ht)} HT")
        if missing:
            z.writestr("A_LIRE_justificatifs_manquants.txt", "Justificatifs manquants\r\n\r\n" + "\r\n".join(missing))
    return buf.getvalue(), base + ".zip", "application/zip"


# ---------------------------------------------------------------------------
# Rapport PDF direction (SPEC §8.2) — reportlab
# ---------------------------------------------------------------------------
def ceo_report_pdf(e: BillingEngine, logo_path: str | None = None) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    BRAND, TEAL, ORANGE, RED, INK, MUTED, LINE, BG = (
        colors.HexColor(c)
        for c in ["#6A3FBF", "#24969A", "#B9741A", "#B93A27", "#1E1846", "#625E80", "#E3E0EF", "#F5F4FA"]
    )
    d = e.dashboard()
    cash = d["cash"]
    month = month_label(e.now)
    year = e.now[:4]
    H1 = ParagraphStyle("h1", fontName="Helvetica-Bold", fontSize=18, textColor=INK, spaceAfter=4)
    H2 = ParagraphStyle("h2", fontName="Helvetica-Bold", fontSize=12, textColor=INK, spaceBefore=10, spaceAfter=2)
    SUB = ParagraphStyle("sub", fontName="Helvetica", fontSize=8.5, textColor=MUTED, spaceAfter=6)
    CELL = ParagraphStyle("cell", fontName="Helvetica", fontSize=8.5, textColor=INK, leading=10.5)

    def tbl(head, rows, widths, aligns, total=False):
        data = [head] + [
            [Paragraph(str(c), CELL) if isinstance(c, str) and len(c) > 18 else c for c in r] for r in rows
        ]
        t = Table(data, colWidths=[w * mm for w in widths], repeatRows=1)
        st = [
            ("BACKGROUND", (0, 0), (-1, 0), BG),
            ("TEXTCOLOR", (0, 0), (-1, 0), MUTED),
            ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 8),
            ("FONT", (0, 1), (-1, -1), "Helvetica", 8.5),
            ("TEXTCOLOR", (0, 1), (-1, -1), INK),
            ("LINEBELOW", (0, 1), (-1, -1), 0.3, LINE),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
        ]
        for i, a in enumerate(aligns):
            if a == "r":
                st.append(("ALIGN", (i, 0), (i, -1), "RIGHT"))
        if total:
            st += [("FONT", (0, -1), (-1, -1), "Helvetica-Bold", 8.5), ("LINEABOVE", (0, -1), (-1, -1), 0.6, INK)]
        t.setStyle(TableStyle(st))
        return t

    def kpis(items):
        row = [
            [
                Paragraph(
                    f"<font size=12 color='{c}'><b>{v}</b></font><br/><font size=7.5 color='#625E80'>{lbl}</font>", CELL
                )
                for v, lbl, c in items
            ]
        ]
        t = Table(row, colWidths=[(178 / len(items)) * mm] * len(items))
        t.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, -1), BG),
                    ("BOX", (0, 0), (-1, -1), 0, BG),
                    ("TOPPADDING", (0, 0), (-1, -1), 6),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
                ]
            )
        )
        return t

    def header_footer(canvas, doc):
        canvas.saveState()
        canvas.setFillColor(BRAND)
        canvas.rect(0, A4[1] - 22 * mm, A4[0], 22 * mm, fill=1, stroke=0)
        if logo_path:
            with contextlib.suppress(Exception):
                canvas.drawImage(logo_path, 16 * mm, A4[1] - 18 * mm, 14 * mm, 14 * mm, mask="auto")
        canvas.setFillColor(colors.white)
        canvas.setFont("Helvetica-Bold", 15)
        canvas.drawString(34 * mm, A4[1] - 11 * mm, "DIGITALIA")
        canvas.setFont("Helvetica", 9.5)
        canvas.drawString(34 * mm, A4[1] - 17 * mm, f"Rapport de facturation · {month}")
        canvas.setFont("Helvetica", 8.5)
        canvas.drawRightString(A4[0] - 16 * mm, A4[1] - 11 * mm, f"Édité le {e.today:%d/%m/%Y}")
        canvas.drawRightString(A4[0] - 16 * mm, A4[1] - 17 * mm, "Confidentiel")
        canvas.setStrokeColor(LINE)
        canvas.line(16 * mm, 10 * mm, A4[0] - 16 * mm, 10 * mm)
        canvas.setFillColor(MUTED)
        canvas.setFont("Helvetica", 8)
        canvas.drawString(16 * mm, 6 * mm, f"DIGITALIA · Facturation clients · {month}")
        canvas.drawRightString(A4[0] - 16 * mm, 6 * mm, f"Page {doc.page}")
        canvas.restoreState()

    prev = d["prev_month_revenue"]
    rev = d["month_revenue"]
    pct = f" · {'+' if rev >= prev else ''}{round((rev - prev) / prev * 100)} % vs mois précédent" if prev else ""
    story = [
        Paragraph(f"Facturation · {month}", H1),
        Paragraph("Synthèse de l'activité, des encaissements attendus et du carnet de commandes.", SUB),
        kpis(
            [
                (eur(rev, 0), f"Facturé en {month}{pct}", "#6A3FBF"),
                (eur(d["ytd"], 0), f"Cumul facturé {year}", "#1E1846"),
                (eur(d["mrr"], 0), "Récurrent fixe par mois", "#24969A"),
                (eur(d["to_bill_total"], 0), "Encore à facturer ce mois", "#B9741A"),
            ]
        ),
        Spacer(1, 6),
    ]
    done, due = d["billed_this_month_total"], d["to_bill_total"]
    done_pct = round(done / (done + due) * 100) if done + due else 0
    story += [
        Paragraph(f"Facturation de {month}", H2),
        Paragraph(
            f"{eur(done + due, 0)} prévus · {eur(done, 0)} déjà facturés ({done_pct} %) · "
            f"{eur(due, 0)} encore à facturer",
            SUB,
        ),
    ]
    rows = [
        ["Facturé", e.client(x.project.client_id).name, x.project.name, x.label, eur(x.amount)]
        for x in d["billed_this_month"]
    ] + [
        ["À facturer", e.client(x.project.client_id).name, x.project.name, x.label, eur(x.amount)] for x in d["to_bill"]
    ]
    if rows:
        story.append(
            tbl(
                ["Statut", "Client", "Projet", "Détail", "Montant"],
                rows + [["", "", "", "Total du mois", eur(done + due)]],
                [20, 30, 46, 58, 24],
                "lllr" + "r",
                total=True,
            )
        )
    story += [
        Paragraph(f"Trésorerie · à encaisser d'ici fin {month.split()[0]}", H2),
        Paragraph(
            f"Factures émises non payées : {eur(cash['expected_by_month_end'], 0)} attendus avant la fin du mois · "
            f"{eur(cash['later_total'], 0)} sur les mois suivants",
            SUB,
        ),
        kpis(
            [
                (eur(cash["late_total"], 0), f"En retard · {len(cash['late'])} facture(s)", "#B93A27"),
                (
                    eur(cash["this_month_total"], 0),
                    f"Échéance ce mois · {len(cash['this_month'])} facture(s)",
                    "#B9741A",
                ),
                (eur(cash["later_total"], 0), "Mois suivants", "#1E1846"),
            ]
        ),
        Spacer(1, 4),
    ]
    allc = cash["late"] + cash["this_month"] + cash["no_due"] + [x for v in cash["later_by_month"].values() for x in v]
    if allc:
        story.append(
            tbl(
                ["Client", "Projet", "Facture", "Échéance", "Montant"],
                [
                    [
                        e.client(x.project.client_id).name,
                        x.project.name,
                        x.label,
                        ("En retard" + (f" ({x.due_date:%d/%m})" if x.due_date else ""))
                        if x.pay_state == "late"
                        else (f"{x.due_date:%d/%m/%Y}" if x.due_date else "Non renseignée"),
                        eur(x.amount),
                    ]
                    for x in allc
                ]
                + [["", "", "", "Total", eur(cash["total"])]],
                [30, 40, 58, 26, 24],
                "llllr",
                total=True,
            )
        )
    story.append(Paragraph("Carnet de commandes · reste à facturer", H2))
    if d["remaining"]:
        story += [
            Paragraph(f"{eur(sum(r['remaining'] for r in d['remaining']), 0)} sur les projets validés", SUB),
            tbl(
                ["Client", "Projet", "Devis / contrat", "Facturé", "Reste", "Planifié"],
                [
                    [
                        e.client(r["project"].client_id).name,
                        r["project"].name,
                        eur(r["total"], 0),
                        eur(r["billed"], 0),
                        eur(r["remaining"], 0),
                        eur(r["planned"], 0) if r["planned"] else "—",
                    ]
                    for r in d["remaining"]
                ],
                [34, 56, 24, 22, 22, 20],
                "llrrrr",
            ),
        ]
    if d["proposals"]:
        story += [
            Paragraph("Propositions en attente de validation", H2),
            tbl(
                ["Client", "Projet", "Envoyée", "Montant"],
                [
                    [
                        e.client(r["project"].client_id).name,
                        r["project"].name + (" — " + r["project"].note if r["project"].note else ""),
                        month_label(r["project"].start_month),
                        eur(r["value"], 0),
                    ]
                    for r in d["proposals"]
                ],
                [34, 100, 22, 22],
                "lllr",
            ),
        ]
    months = month_range(year + "-01", e.now)
    byc = sorted(
        [
            (
                c,
                sum(
                    (e.project_month(p, m).revenue for m in months for p in e.s.projects if p.client_id == c.id),
                    Decimal(0),
                ),
            )
            for c in e.s.clients
        ],
        key=lambda x: -x[1],
    )
    byc = [x for x in byc if x[1]]
    story += [
        Paragraph(f"Chiffre d'affaires {year} par client", H2),
        tbl(
            ["Client", "Part", f"Facturé {year}"],
            [[c.name, f"{round(v / d['ytd'] * 100) if d['ytd'] else 0} %", eur(v, 0)] for c, v in byc]
            + [["Total", "100 %", eur(d["ytd"], 0)]],
            [110, 30, 38],
            "lrr",
            total=True,
        ),
        Paragraph("Projets les plus rentables", H2),
        Paragraph("12 derniers mois", SUB),
        tbl(
            ["#", "Projet", "Client", "Type", "CA 12 mois"],
            [
                [
                    i + 1,
                    r["project"].name,
                    e.client(r["project"].client_id).name,
                    "Fixe" if r["project"].is_fixed else "Devis",
                    eur(r["revenue"], 0),
                ]
                for i, r in enumerate(d["ranking_12"][:8])
            ],
            [8, 76, 46, 18, 30],
            "lllllr",
        ),
    ]
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf, pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm, topMargin=30 * mm, bottomMargin=16 * mm
    )
    doc.build(story, onFirstPage=header_footer, onLaterPages=header_footer)
    return buf.getvalue()
